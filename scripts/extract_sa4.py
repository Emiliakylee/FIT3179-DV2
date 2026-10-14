"""Build the SA4 spike map data (C9).

Usage:
    python3 scripts/extract_sa4.py path/to/SA4__2025_Dec_.xlsx

Joins international higher education enrolments by SA4 to SA4 label points
generated from the ABS boundary file with:
    mapshaper data/au_sa4.topojson -points inner -o format=geojson

Sources: Department of Education, International Student Enrolments and
Commencements by SA4 Location (year to date December 2025, PRISMS);
ABS ASGS Edition 3 SA4 boundaries.
"""
import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(sys.argv[1])
POINTS = Path(sys.argv[2] if len(sys.argv) > 2 else "/tmp/sa4_pts.json")
OUT = Path(__file__).resolve().parent.parent / "data"

# The Department's file misspells one ABS region name.
ALIAS = {"Warnambool and South West": "Warrnambool and South West"}

wb = load_workbook(SRC, read_only=True, data_only=True)
ws = wb["Data"]

totals = {}
for row in ws.iter_rows(min_row=7, values_only=True):
    state, sa4, _remote, sector, _field, enrolments = row[:6]
    if sector != "Higher Education" or not sa4:
        continue
    try:
        n = float(str(enrolments).replace(",", ""))
    except (TypeError, ValueError):
        continue
    key = ALIAS.get(str(sa4).strip(), str(sa4).strip())
    rec = totals.setdefault(key, {"sa4_name": key, "state": state,
                                  "enrolments": 0.0})
    rec["enrolments"] += n

points = {}
for f in json.load(open(POINTS))["features"]:
    lon, lat = f["geometry"]["coordinates"]
    points[f["properties"]["SA4_NAME21"].strip()] = (lon, lat)

rows, missing = [], []
for name, rec in totals.items():
    if name not in points:
        missing.append(name)
        continue
    lon, lat = points[name]
    rows.append({"sa4_name": name, "state": rec["state"],
                 "enrolments": int(rec["enrolments"]),
                 "lon": round(lon, 4), "lat": round(lat, 4)})
rows.sort(key=lambda r: -r["enrolments"])

with open(OUT / "international_by_sa4_2025.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["sa4_name", "state", "enrolments",
                                      "lon", "lat"])
    w.writeheader()
    w.writerows(rows)

print(f"SA4 regions written : {len(rows)}")
if missing:
    print(f"no boundary match   : {missing}")
print(f"total enrolments    : {sum(r['enrolments'] for r in rows):,}")
for r in rows[:8]:
    print(f"   {r['sa4_name']:<34} {r['enrolments']:>7,}")

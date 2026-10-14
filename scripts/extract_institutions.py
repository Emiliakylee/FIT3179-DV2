"""Build the state aggregate (C7) and the spread comparison (C12).

Usage:
    python3 scripts/extract_institutions.py SES_2025_National_Report_Tables.xlsx

Institution names are used only to map each university to a state and to
weight the state averages. They are never written to the output files, so
nothing in the repository can be read as a university ranking.

Source: QILT, 2025 Student Experience Survey National Report Tables,
sheets FOCUS_UG_UNI_1Y_INST_CI, FOCUS_PGC_UNI_1Y_INST_CI and
RR_ALL_UNI_17-YY_INST.
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(sys.argv[1] if len(sys.argv) > 1
           else "SES_2025_National_Report_Tables.xlsx")
OUT = Path(__file__).resolve().parent.parent / "data"
OUT.mkdir(exist_ok=True)

LEVELS = {"UG": "Undergraduate", "PGC": "Postgraduate"}

# Each university's state comes from the Department of Education's own
# classification (Table 2.5), not from a judgement of ours. Australian
# Catholic University is classified Multi-State there and is therefore left
# out of the state averages, as is Australian University of Theology, which
# the Department's table does not list separately.
DOE_TABLE = Path("/mnt/user-data/uploads/2024_Section2_All_Students.xlsx")

ALIAS = {
    "CQUniversity": "Central Queensland University",
    "Torrens University Australia": "Torrens University",
    "University of South Australia": "The University of South Australia",
    "The University of New England": "University of New England",
    "The University of Newcastle": "University of Newcastle",
}
EXCLUDE = {"Australian Catholic University", "Australian University of Theology"}

STATE_NAME = {
    "NSW": "New South Wales", "VIC": "Victoria", "QLD": "Queensland",
    "WA": "Western Australia", "SA": "South Australia", "TAS": "Tasmania",
    "ACT": "Australian Capital Territory", "NT": "Northern Territory",
}


def tidy(name):
    """Drop footnote markers the two sources attach to institution names."""
    name = re.sub(r"[*†]+$", "", str(name)).strip()
    return re.sub(r"\(\d+\.\d+\)\s*$", "", name).strip()


def parse_ci(cell):
    """'84.5 (83.9, 85.1)' -> (84.5, 83.9, 85.1)"""
    if cell is None:
        return None
    nums = re.findall(r"-?\d+\.?\d*", str(cell))
    if len(nums) < 3:
        return None
    return tuple(float(n) for n in nums[:3])


def clean_measure(text):
    return re.sub(r"^(Focus areas|Questionnaire item):\s*", "",
                  str(text)).replace("*", "").strip()


def build_state_map():
    """institution name -> state abbreviation, from Department Table 2.5."""
    abbrev = {v: k for k, v in STATE_NAME.items()}
    wb_doe = load_workbook(DOE_TABLE, read_only=True, data_only=True)
    mapping = {}
    for row in wb_doe["2.5"].iter_rows(min_row=4, values_only=True):
        state, inst = row[0], row[1]
        if not inst or state not in abbrev:
            continue
        raw = re.sub(r"\(\d+\.\d+\)\s*$", "", str(inst)).strip()
        name = ALIAS.get(raw, raw)
        mapping[name] = abbrev[state]
    return mapping


STATE = build_state_map()

wb = load_workbook(SRC, read_only=True, data_only=True)

# Responses per university in 2025 (last N column of the response-rate sheet).
responses = {}
rr = list(wb["RR_ALL_UNI_17-YY_INST"].iter_rows(values_only=True))
year_row = rr[2]
col_2025 = max(i for i, v in enumerate(year_row) if str(v).strip() == "2025")
for row in rr[4:]:
    if not row[1]:
        continue
    name = tidy(row[1])
    raw = row[col_2025 - 1] if col_2025 else None
    try:
        responses[name] = float(str(raw).replace(",", ""))
    except (TypeError, ValueError):
        responses[name] = 0.0

records = []   # one per institution x level x measure
skipped = set()
for key, level in LEVELS.items():
    rows = list(wb[f"FOCUS_{key}_UNI_1Y_INST_CI"].iter_rows(values_only=True))
    measures = [clean_measure(c) for c in rows[3][2:8]]
    for row in rows[4:]:
        if not row[1]:
            continue
        name = tidy(row[1])
        if name == "All Universities" or name in EXCLUDE:
            continue
        if name not in STATE:
            skipped.add(name)
            continue
        for measure, cell in zip(measures, row[2:8]):
            parsed = parse_ci(cell)
            if parsed is None:
                continue
            value, lo, hi = parsed
            records.append({"institution": name, "state": STATE[name],
                            "level": level, "measure": measure,
                            "value": value, "lo": lo, "hi": hi,
                            "weight": responses.get(name, 0.0)})

# ------------------------------------------------- state weighted averages
state_rows = []
groups = {}
for r in records:
    groups.setdefault((r["state"], r["level"], r["measure"]), []).append(r)
for (state, level, measure), rs in sorted(groups.items()):
    weights = sum(r["weight"] for r in rs)
    if weights == 0:
        continue
    weighted = sum(r["value"] * r["weight"] for r in rs) / weights
    state_rows.append({
        "state": state,
        "state_name": STATE_NAME[state],
        "level": level,
        "measure": measure,
        "value": round(weighted, 1),
        "universities": len(rs),
        "responses": int(weights),
    })

with open(OUT / "ses_state_2025.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["state", "state_name", "level",
                                      "measure", "value", "universities",
                                      "responses"])
    w.writeheader()
    w.writerows(state_rows)

# ------------------------------------- spread comparison, overall only (C12)
MEASURE = "Overall Educational Experience"
spread = []

for i, r in enumerate(sorted((r for r in records if r["measure"] == MEASURE),
                             key=lambda r: (r["level"], r["value"]))):
    spread.append({"lens": "Universities", "level": r["level"],
                   "item": f"University {i % 43 + 1:02d}", "value": r["value"]})

for r in state_rows:
    if r["measure"] == MEASURE:
        spread.append({"lens": "States and territories", "level": r["level"],
                       "item": r["state"], "value": r["value"]})

with open(OUT / "ses_study_area_2025.csv") as f:
    for row in csv.DictReader(f):
        if row["measure"] == MEASURE:
            spread.append({"lens": "Study areas", "level": row["level"],
                           "item": row["study_area"],
                           "value": float(row["value"])})

with open(OUT / "ses_spread_2025.csv", "w", newline="") as f:
    w = csv.DictWriter(f, fieldnames=["lens", "level", "item", "value"])
    w.writeheader()
    w.writerows(spread)

print(f"institutions parsed : {len(set(r['institution'] for r in records))}")
if skipped:
    print(f"not matched to a state: {sorted(skipped)}")
print(f"state rows          : {len(state_rows)}")
print(f"spread rows         : {len(spread)}")
